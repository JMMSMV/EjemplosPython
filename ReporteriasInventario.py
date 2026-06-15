import pandas as pd
import tkinter as tk
from tkinter import ttk
from openpyxl import load_workbook

# --- Funciones de reporte ---
def reporte_por_dia(df, tc):
    df["Fecha"] = pd.to_datetime(df["Fecha"], dayfirst=True)
    pivot = pd.pivot_table(
        df,
        values="Cantidad",
        index="Descripcion",
        columns=df["Fecha"].dt.strftime("%d/%m/%Y"),
        aggfunc="sum",
        fill_value=0
    )
    pivot["Total"] = pivot.sum(axis=1)
    pivot.loc["TOTAL GENERAL"] = pivot.sum()
    ruta1= "/storage/emulated/0/Reporte_por_dia.xlsx"
    # ruta = "Reporte_por_dia.xlsx"
    pivot.to_excel(ruta1)
    return ruta1,pivot

def reporte_compacto(df, tc):
    df["Importe"] = df["Cantidad"] * df["Precio"]
    tabla = df.groupby(["Codigo","Descripcion"]).agg(
        Cantidad=("Cantidad","sum"),
        Precio=("Precio","mean"),
        Importe_Total=("Importe","sum"),
        Importe_Transferencia=("Importe", lambda x: df.loc[x.index][df.loc[x.index,"TipoPago"]=="Transferencia"]["Importe"].sum()),
        Importe_Efectivo=("Importe", lambda x: df.loc[x.index][df.loc[x.index,"TipoPago"]=="Efectivo"]["Importe"].sum())
    ).reset_index()

    total_usd = tabla["Importe_Total"].sum()
    total_cup = total_usd * tc

    ruta2= "/storage/emulated/0/Reporte_compacto.xlsx"
    # ruta = "Reporte_compacto.xlsx"
    tabla.to_excel(ruta2, index=False)

    wb = load_workbook(ruta2)
    ws = wb.active
    fila_final = ws.max_row + 2
    ws.cell(row=fila_final, column=1, value="IMPORTE EN USD")
    ws.cell(row=fila_final, column=2, value=total_usd)
    ws.cell(row=fila_final+1, column=1, value="TASA DE CAMBIO")
    ws.cell(row=fila_final+1, column=2, value=tc)
    ws.cell(row=fila_final+2, column=1, value="IMPORTE EN CUP")
    ws.cell(row=fila_final+2, column=2, value=total_cup)
    wb.save(ruta2)
    return ruta2, tabla

# --- Interfaz gráfica ---
def generar_reporte():
    opcion = combo.get()
    tc = float(entry_tc.get())
    ruta3= "/storage/emulated/0/Seleccionados.xlsx"
    df = pd.read_excel(ruta3)

    if opcion == "Reporte por día":
        archivo, tabla = reporte_por_dia(df, tc)
    elif opcion == "Reporte compacto":
        archivo, tabla = reporte_compacto(df, tc)
    else:
        archivo, tabla = None, None

    if archivo:
        lbl_result.config(text=f"Reporte generado: {archivo}")
        # Mostrar tabla en pantalla
        text_box.delete("1.0", tk.END)
        text_box.insert(tk.END, str(tabla))

# Ventana principal
root = tk.Tk()
root.title("Generador de Reportes")

# Botón salir arriba a la derecha
btn_salir = tk.Button(root, text="Salir aplicación", command=root.quit, bg="red", fg="white")
btn_salir.pack(anchor="ne", padx=10, pady=5)

tk.Label(root, text="Seleccione reporte:").pack(pady=5)
combo = ttk.Combobox(root, values=["Reporte por día","Reporte compacto"])
combo.current(0)
combo.pack(pady=5)

tk.Label(root, text="Tasa de cambio (tc):").pack(pady=5)
entry_tc = tk.Entry(root)
entry_tc.insert(0,"545.00")  # valor por defecto
entry_tc.pack(pady=5)

btn = tk.Button(root, text="Generar", command=generar_reporte)
btn.pack(pady=10)

lbl_result = tk.Label(root, text="")
lbl_result.pack(pady=5)

# Caja de texto para mostrar reporte
text_box = tk.Text(root, width=100, height=30)
text_box.pack(pady=10)

root.mainloop()